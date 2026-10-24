import io
import json
import logging
import re
import time
from typing import Any

from app.core.config import get_settings

logger = logging.getLogger(__name__)


def extract_text_from_bytes(content: bytes, filename: str) -> str:
    """Extracts text from PDF, DOCX, or plain text bytes."""
    fname_lower = filename.lower()
    extracted_text = ""

    if fname_lower.endswith(".pdf"):
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            extracted_text = "\n".join(pages)
        except Exception as e:
            logger.warning(f"pypdf extraction failed on {filename}: {e}")

    elif fname_lower.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            extracted_text = "\n".join(paragraphs)
        except Exception as e:
            logger.warning(f"python-docx extraction failed on {filename}: {e}")

    # Fallback or plain text file
    if not extracted_text:
        try:
            extracted_text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                extracted_text = content.decode("latin-1")
            except Exception:
                extracted_text = ""

    # Normalize whitespace
    extracted_text = re.sub(r"\r\n", "\n", extracted_text)
    extracted_text = re.sub(r"[ \t]+", " ", extracted_text)
    extracted_text = re.sub(r"\n{3,}", "\n\n", extracted_text).strip()

    return extracted_text


def _clean_json_text(text: str) -> str:
    text = text.strip()
    match = re.search(r"```(?:json)?\s*([\s\S]*?)\s*```", text)
    if match:
        return match.group(1).strip()
    return text


def _heuristic_resume_analysis(raw_text: str, target_role: str) -> dict[str, Any]:
    """Reliable heuristic fallback when Groq LLM is unavailable."""
    text_lower = raw_text.lower()

    # Detect skills
    all_languages = ["python", "java", "c++", "c", "javascript", "typescript", "go", "rust", "php", "ruby", "kotlin", "swift", "sql"]
    all_frameworks = ["react", "node", "express", "fastapi", "django", "flask", "spring boot", "next.js", "vue", "angular", "tailwind", "pytorch", "tensorflow", "scikit-learn"]
    all_databases = ["postgresql", "postgres", "mysql", "mongodb", "sqlite", "redis", "cassandra", "dynamodb"]
    all_tools = ["git", "github", "docker", "kubernetes", "aws", "gcp", "azure", "linux", "postman", "jira", "ci/cd"]

    detected_langs = [l.title() for l in all_languages if re.search(rf"\b{re.escape(l)}\b", text_lower)]
    detected_frameworks = [f.title() for f in all_frameworks if re.search(rf"\b{re.escape(f)}\b", text_lower)]
    detected_dbs = [d.title() for d in all_databases if re.search(rf"\b{re.escape(d)}\b", text_lower)]
    detected_tools = [t.upper() if len(t) <= 4 else t.title() for t in all_tools if re.search(rf"\b{re.escape(t)}\b", text_lower)]

    # Attempt to extract project headers
    lines = raw_text.split("\n")
    projects = []
    current_proj = None
    for line in lines:
        line_clean = line.strip()
        if not line_clean:
            continue
        if any(kw in line_clean.lower() for kw in ["project", "github.com", "app", "system", "portal", "clone", "platform"]):
            if len(line_clean) < 60 and not line_clean.lower().startswith("projects"):
                projects.append({
                    "title": line_clean[:60],
                    "technologies": [tech for tech in (detected_langs + detected_frameworks) if tech.lower() in line_clean.lower()],
                    "summary": "Technical implementation and delivery documented on candidate resume.",
                    "claimed_impact": "Demonstrates practical engineering problem solving."
                })
        if len(projects) >= 4:
            break

    if not projects:
        projects = [
            {
                "title": f"{detected_frameworks[0] if detected_frameworks else 'Full-Stack'} Application",
                "technologies": (detected_langs[:2] + detected_frameworks[:2]) or ["Python", "FastAPI"],
                "summary": "Core software project highlighted in candidate resume.",
                "claimed_impact": "Implemented functional architecture with database connectivity.",
            }
        ]

    # ATS score estimation
    score = 65
    if len(detected_langs) >= 2:
        score += 8
    if len(detected_frameworks) >= 2:
        score += 8
    if len(detected_dbs) >= 1:
        score += 7
    if len(detected_tools) >= 2:
        score += 7
    score = min(92, max(60, score))

    # Generate tailored interview questions from detected items
    generated_questions = []

    # Q1: Primary Project Architecture
    first_proj = projects[0]
    p_tech = ", ".join(first_proj.get("technologies") or detected_frameworks[:2] or ["the chosen stack"])
    generated_questions.append({
        "id": "res-q1",
        "question": f"In your project '{first_proj['title']}', why did you select {p_tech}? Walk me through your high-level system architecture and how data flows from user input to the database.",
        "category": "Resume & Projects",
        "topic": f"Architecture: {first_proj['title']}",
        "difficulty": "Intermediate",
        "rubric_key": "technical",
        "interviewer_context": f"Interviewers inspect candidate-claimed projects to verify hands-on authorship and architectural reasoning rather than copy-pasted code.",
        "key_points": [
            "Clear separation of concerns (frontend/API/database)",
            "Concrete justification for the technology stack chosen",
            "Data validation and error handling mechanisms",
            "Handling concurrent user requests and database transactions",
        ],
        "ideal_answer_outline": "Explain problem statement, justify stack choices against alternatives, describe components and APIs, and conclude with performance/security measures.",
    })

    # Q2: Challenging Bug / Deep Technical Hurdle
    generated_questions.append({
        "id": "res-q2",
        "question": f"What was the most challenging technical roadblock or bug you encountered while building '{first_proj['title']}', and how did you diagnose and resolve it?",
        "category": "Resume & Projects",
        "topic": "Debugging & Resilience",
        "difficulty": "Intermediate",
        "rubric_key": "technical",
        "interviewer_context": "Probes debugging methodology, root-cause analysis, and whether the candidate understands their stack under failure conditions.",
        "key_points": [
            "Specific description of the symptom vs the underlying root cause",
            "Tools or profiling techniques used (logs, breakpoints, network tab, explain query)",
            "The permanent architectural or code fix implemented",
            "Key takeaway or testing added to prevent recurrence",
        ],
        "ideal_answer_outline": "Set context of the issue, describe isolation process, show structured debugging, explain the fix and unit test added.",
    })

    # Q3: Technical Skill Verification (Language / Framework)
    primary_lang = detected_langs[0] if detected_langs else "Python"
    generated_questions.append({
        "id": "res-q3",
        "question": f"Your resume highlights strong proficiency in {primary_lang}. How does {primary_lang} manage memory, asynchronous execution, and thread safety under heavy workloads?",
        "category": "Technical",
        "topic": f"Core {primary_lang} Internals",
        "difficulty": "Intermediate",
        "rubric_key": "technical",
        "interviewer_context": f"Validates whether candidate understands underlying runtime internals or only surface-level syntax.",
        "key_points": [
            f"Core runtime execution model of {primary_lang}",
            "Memory management (garbage collection, reference counting, stack vs heap)",
            "Concurrency primitives (asyncio, thread pools, event loop)",
        ],
        "ideal_answer_outline": f"Articulate memory management lifecycle, concurrency mechanisms, and best practices for writing production-ready {primary_lang}.",
    })

    # Q4: Database & Schema Design
    db_name = detected_dbs[0] if detected_dbs else "PostgreSQL"
    generated_questions.append({
        "id": "res-q4",
        "question": f"You mentioned using {db_name} in your projects. How did you structure your schema, what indexing strategies did you employ, and how would you optimize a slow query?",
        "category": "SQL & Analytics",
        "topic": f"Database Design & Optimization ({db_name})",
        "difficulty": "Intermediate",
        "rubric_key": "technical",
        "interviewer_context": "Evaluates database schema normalization, indexing trade-offs, and query performance debugging.",
        "key_points": [
            "Normalization vs deliberate denormalization",
            "B-tree index usage and compound index column ordering",
            "EXPLAIN ANALYZE for diagnosing query execution bottlenecks",
        ],
        "ideal_answer_outline": "Explain relational structure, index selection rationale, execution plan profiling, and connection pooling.",
    })

    # Q5: Behavioral / Engineering Ownership
    generated_questions.append({
        "id": "res-q5",
        "question": "Tell me about a time when you had to learn an unfamiliar library, framework, or API on a tight deadline for a project. How did you break it down?",
        "category": "Behavioral",
        "topic": "Rapid Learning & Agility",
        "difficulty": "Beginner",
        "rubric_key": "hr",
        "interviewer_context": "Assesses learning agility, resourcefulness, and self-direction which campus interviewers prioritize.",
        "key_points": [
            "Clear context and initial knowledge gap",
            "Structured approach to research and documentation",
            "Building small proof-of-concepts before integration",
            "Delivering on schedule with measurable results",
        ],
        "ideal_answer_outline": "Use STAR format: describe situation, task, methodical learning action, and positive project outcome.",
    })

    # Q6: Trade-offs & Future Scaling
    generated_questions.append({
        "id": "res-q6",
        "question": f"If the user traffic for your project '{first_proj['title']}' scaled from a few users to 100,000 daily active users, what would break first and what architectural changes would you implement?",
        "category": "Resume & Projects",
        "topic": "Scalability & Bottlenecks",
        "difficulty": "Advanced",
        "rubric_key": "technical",
        "interviewer_context": "Tests system design intuition, caching layers, and database load mitigation.",
        "key_points": [
            "Identification of single points of failure (database connection limits, CPU-heavy tasks)",
            "Caching strategy (Redis/Memcached) for frequent reads",
            "Asynchronous task queues (Celery/RabbitMQ) for background jobs",
            "Horizontal scaling behind a load balancer",
        ],
        "ideal_answer_outline": "Identify bottlenecks systematically (DB, state, network), propose caching and queuing, and explain horizontal scaling.",
    })

    return {
        "parsed_data": {
            "candidate_name": "Candidate",
            "skills": {
                "languages": detected_langs,
                "frameworks_and_libraries": detected_frameworks,
                "databases": detected_dbs,
                "tools_and_platforms": detected_tools,
                "core_competencies": ["Data Structures", "System Design Fundamentals", "Object-Oriented Programming"],
            },
            "projects": projects,
            "experience": [],
            "education": [],
        },
        "analysis": {
            "overview_summary": f"Resume indicates solid foundational hands-on experience targeting {target_role}. Features projects utilizing {', '.join(detected_langs[:2] + detected_frameworks[:2]) or 'modern developer technologies'}.",
            "ats_score": score,
            "ats_breakdown": {
                "skill_relevance": min(95, score + 4),
                "project_depth": min(90, score - 2),
                "impact_metrics": max(50, score - 15),
                "formatting_structure": 82,
            },
            "strengths": [
                f"Demonstrated multi-tier project development using {', '.join((detected_langs[:2] + detected_frameworks[:2]) or ['modern tools'])}.",
                f"Practical exposure to database technologies ({', '.join(detected_dbs) or 'SQL databases'}).",
                "Clear project descriptions suitable for technical defense.",
            ],
            "flagged_risks": [
                "Quantifiable business metrics (e.g., latency reduction, user count, test coverage %) can be made more explicit in project bullet points.",
                "Be prepared for deep-dive questioning on internal runtime mechanisms of listed languages.",
                "Ensure you can clearly justify why you chose your specific database over alternative options.",
            ],
            "recommended_focus_areas": [
                f"Deepen knowledge of {primary_lang} memory management and asynchronous primitives.",
                "Practice explaining end-to-end data flow and architectural trade-offs using the STAR method.",
                "Review SQL query optimization and indexing strategies.",
            ],
            "generated_questions": generated_questions,
        },
    }


def analyze_resume_and_generate_questions(
    raw_text: str,
    target_role: str = "Software Development Engineer",
) -> dict[str, Any]:
    """Uses Groq LLM to analyze candidate resume, extract structured data, and generate tailored questions."""
    settings = get_settings()

    if not raw_text or len(raw_text.strip()) < 50:
        raise ValueError("Resume text is too short or empty to analyze.")

    # Truncate text to avoid exceeding token limits while preserving essential content
    clean_text = raw_text[:12000]

    if not settings.groq_api_key or settings.groq_api_key.strip() in {"", "dummy", "mock"}:
        logger.info("Groq API key not set; using heuristic resume analyzer.")
        return _heuristic_resume_analysis(clean_text, target_role)

    start_time = time.perf_counter()
    system_prompt = (
        "You are an elite Senior Technical Recruiter and Engineering Hiring Manager at top tech firms (Google, Microsoft, Amazon, TCS Digital).\n"
        "Your mission is to meticulously analyze a candidate's resume, extract their declared skills and projects, "
        "audit their interview preparedness (ATS & depth), identify claims that interviewers will grill them on, "
        "and generate 6 to 8 realistic, probing interview questions based DIRECTLY on their specific projects, technologies, and bullet points.\n"
        "Return STRICT valid JSON only. Never wrap outside the JSON format."
    )

    user_prompt = f"""Target Role: {target_role}

CANDIDATE RESUME TEXT:
\"\"\"
{clean_text}
\"\"\"

Analyze this resume and produce a structured JSON object matching this exact schema:
{{
  "parsed_data": {{
    "candidate_name": "<Extracted Name or 'Candidate'>",
    "skills": {{
      "languages": ["<Language 1>", "<Language 2>"],
      "frameworks_and_libraries": ["<Framework 1>", "<Framework 2>"],
      "databases": ["<DB 1>"],
      "tools_and_platforms": ["<Tool 1>"],
      "core_competencies": ["<e.g. OOP, REST APIs, Microservices>"]
    }},
    "projects": [
      {{
        "title": "<Project Name>",
        "technologies": ["<Tech used>"],
        "summary": "<1-2 sentence description>",
        "claimed_impact": "<Impact or features built>"
      }}
    ],
    "experience": [
      {{
        "role": "<Role Title>",
        "company": "<Company Name>",
        "duration": "<Dates/Duration>",
        "summary": "<Brief summary of work>"
      }}
    ],
    "education": [
      {{
        "degree": "<Degree>",
        "institution": "<College/University>",
        "graduation_year": "<Year>",
        "gpa_or_percentage": "<Grade if mentioned>"
      }}
    ]
  }},
  "analysis": {{
    "overview_summary": "<Comprehensive 2-3 sentence assessment of the candidate's profile, depth, and alignment with target role>",
    "ats_score": <Integer from 50 to 95 reflecting technical depth, metrics, and clarity>,
    "ats_breakdown": {{
      "skill_relevance": <Integer 0-100>,
      "project_depth": <Integer 0-100>,
      "impact_metrics": <Integer 0-100>,
      "formatting_structure": <Integer 0-100>
    }},
    "strengths": [
      "<Key concrete strength based on resume 1>",
      "<Key concrete strength based on resume 2>",
      "<Key concrete strength based on resume 3>"
    ],
    "flagged_risks": [
      "<Vulnerable claim or missing detail that an interviewer will probe 1>",
      "<Vulnerable claim or missing detail that an interviewer will probe 2>",
      "<Vulnerable claim or missing detail that an interviewer will probe 3>"
    ],
    "recommended_focus_areas": [
      "<Specific technical topic to revise 1>",
      "<Specific technical topic to revise 2>",
      "<Specific technical topic to revise 3>"
    ],
    "generated_questions": [
      {{
        "id": "res-q1",
        "question": "<High-impact question probing their specific project architecture, decision-making, or declared stack>",
        "category": "<'Resume & Projects' | 'Python' | 'SQL & Analytics' | 'ML, AI & GenAI' | 'Technical' | 'Behavioral'>",
        "topic": "<Project Title or Specific Tech Area>",
        "difficulty": "<'Beginner' | 'Intermediate' | 'Advanced'>",
        "rubric_key": "<'technical' | 'hr'>",
        "interviewer_context": "<Why an interviewer asks this based on what they wrote on their resume>",
        "key_points": [
          "<Key expectation in candidate's answer 1>",
          "<Key expectation in candidate's answer 2>",
          "<Key expectation in candidate's answer 3>"
        ],
        "ideal_answer_outline": "<Guidance on how the candidate should structure their winning response>"
      }}
    ]
  }}
}}
Generate at least 6 thorough, tailored questions (minimum 3 on their projects, 2 on their declared core tech stack/databases, 1 on engineering challenges/behavioral)."""

    try:
        from groq import Groq

        client = Groq(api_key=settings.groq_api_key, timeout=40.0)
        completion = client.chat.completions.create(
            model=settings.groq_model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.3,
            response_format={"type": "json_object"},
        )

        latency_ms = int((time.perf_counter() - start_time) * 1000)
        logger.info(f"Groq resume analysis completed in {latency_ms}ms")
        raw_content = completion.choices[0].message.content or "{}"
        parsed = json.loads(_clean_json_text(raw_content))

        # Basic validation of expected structure
        if "parsed_data" in parsed and "analysis" in parsed and parsed["analysis"].get("generated_questions"):
            return parsed
        else:
            logger.warning("Groq response missed expected schema keys; using fallback.")
            return _heuristic_resume_analysis(clean_text, target_role)

    except Exception as e:
        logger.error(f"Groq resume analysis failed: {e}. Utilizing fallback.")
        return _heuristic_resume_analysis(clean_text, target_role)
